"""
Generate realistic dummy data covering every reconciliation category.

Creates (in ``sample_data/``):
    purchase_register.xlsx        - books (title rows, mixed date formats, total row)
    gstr2b_apr2024.json           - official GSTR-2B JSON layout, return period 04-2024
    gstr2b_may2024.json           - official GSTR-2B JSON layout, return period 05-2024
    gstr2b_portal_format.xlsx     - same 2B data in the portal Excel layout (B2B / B2B-CDNR sheets)
    sales_register.xlsx           - sales books
    gstr1_may2024.json            - GSTR-1 JSON (item-wise layout)
    expected_categories.csv       - what each books/portal row is designed to test
and blank input templates in ``templates/``.

Run:  python sample_data/generate_sample_data.py
"""
from __future__ import annotations

import csv
import json
import sys
from dataclasses import dataclass, field, replace
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from modules.cleaner import gstin_check_char  # noqa: E402

OUT = ROOT / "sample_data"
TEMPLATES = ROOT / "templates"
COMPANY_STATE = "27"
STATE_NAMES = {"27": "Maharashtra", "29": "Karnataka", "24": "Gujarat", "07": "Delhi", "33": "Tamil Nadu",
               "06": "Haryana", "19": "West Bengal", "36": "Telangana", "32": "Kerala", "09": "Uttar Pradesh"}


def make_gstin(state: str, pan: str, entity: str = "1") -> str:
    """Valid GSTIN (correct checksum) from state code and PAN."""
    first14 = f"{state}{pan}{entity}Z"
    return first14 + gstin_check_char(first14)


COMPANY_NAME = "Demo Traders Pvt Ltd"
COMPANY_GSTIN = make_gstin("27", "AABCD1234E")

SUPPLIERS = [  # (state, PAN, name, invoice prefix)
    ("27", "AAACS1234F", "Shree Ganesh Steels Pvt Ltd", "SGS"),
    ("27", "AADFP5678K", "Patil Packaging Industries", "PPI"),
    ("29", "AABCK4321M", "Karnataka Electronics Ltd", "KEL"),
    ("24", "AAGFR8765Q", "Rajkot Engineering Works", "REW"),
    ("07", "AACCD2468L", "Delhi Office Supplies Pvt Ltd", "DOS"),
    ("33", "AAECT1357P", "Chennai Textiles Ltd", "CTL"),
    ("27", "ABCPM9876R", "Mehta Transport Services", "MTS"),
    ("06", "AAHCG1122S", "Gurgaon Logistics Pvt Ltd", "GL"),
    ("19", "AAJCB3344T", "Bengal Chemicals Co", "BCC"),
    ("36", "AAKCH5566U", "Hyderabad Pharma Distributors", "HPD"),
]
SUP = [{"state": s, "gstin": make_gstin(s, p), "name": n, "prefix": x} for s, p, n, x in SUPPLIERS]


@dataclass
class Doc:
    """One document as it appears on one side (books or portal)."""

    gstin: str
    name: str
    inv: str
    dt: date
    taxable: float
    igst: float = 0.0
    cgst: float = 0.0
    sgst: float = 0.0
    cess: float = 0.0
    doc: str = "INV"          # INV / CN / DN
    period: str = ""          # MMYYYY for portal
    itcavl: str = "Y"
    pos: str = ""
    category: str = ""
    extra: dict = field(default_factory=dict)

    @property
    def value(self) -> float:
        return round(self.taxable + self.igst + self.cgst + self.sgst + self.cess, 2)


def taxed(sup: dict, inv: str, dt: date, taxable: float, rate: float = 18.0, doc: str = "INV",
          category: str = "", head: str | None = None) -> Doc:
    """Doc with tax computed from rate; intra-state suppliers get CGST+SGST unless *head* overrides."""
    intra = (head or ("CGST+SGST" if sup["state"] == COMPANY_STATE else "IGST")) == "CGST+SGST"
    tax = round(taxable * rate / 100, 2)
    half = round(tax / 2, 2)
    return Doc(sup["gstin"], sup["name"], inv, dt, round(taxable, 2),
               igst=0.0 if intra else tax, cgst=half if intra else 0.0, sgst=half if intra else 0.0,
               doc=doc, pos=COMPANY_STATE, category=category, period=dt.strftime("%m%Y"))


def build_purchase_data() -> tuple[list[Doc], list[Doc]]:
    """Return (books, portal) documents for the purchase reconciliation demo."""
    books: list[Doc] = []
    portal: list[Doc] = []
    d0 = date(2024, 4, 1)

    # 1. Exact matches (26) - including invoice-number formatting variations
    variants = [
        ("{p}/24-25/{n:04d}", "{p}/{n:04d}"), ("{p}-{n}", "{p}{n}"), ("{p} {n:05d}", "{p}-{n}"),
        ("{p}/{n}/2024-25", "{p}/{n}"), ("{p}#{n}", "{p}.{n}"), ("{p}/{n}", "{p}/{n}"),
    ]
    rates = [18, 12, 5, 28, 18]
    for i in range(26):
        sup = SUP[i % len(SUP)]
        n = 100 + i * 7
        bfmt, pfmt = variants[i % len(variants)]
        dt = d0 + timedelta(days=(i * 2) % 58)
        taxable = 10_000 + (i * 37_531) % 4_90_000 + 0.5 * (i % 2)
        rate = rates[i % len(rates)]
        b = taxed(sup, bfmt.format(p=sup["prefix"], n=n), dt, taxable, rate, category="Exact Match")
        p = replace(b, inv=pfmt.format(p=sup["prefix"], n=n))
        if i in (3, 11, 19):  # April invoices uploaded late by the supplier -> May 2B
            b = replace(b, dt=date(2024, 4, 20 + i % 9), category="Exact Match (cross-period)")
            p = replace(b, inv=p.inv, period="052024", category="Exact Match (cross-period)")
        books.append(b)
        portal.append(p)
    # duplicate entry in books of the first exact invoice
    books.append(replace(books[0], category="Duplicate in books -> Missing in 2B"))

    # 2. Matched with tolerance (5) - rounding differences within ±₹10
    for i, delta in enumerate([1.0, -4.5, 7.25, -9.0, 3.0]):
        sup = SUP[(i + 2) % len(SUP)]
        b = taxed(sup, f"{sup['prefix']}/T{500 + i}", d0 + timedelta(days=5 + i * 9), 25_000 + i * 11_000,
                  category="Matched with Tolerance")
        p = replace(b, taxable=round(b.taxable + delta, 2))
        if p.igst:
            p = replace(p, igst=round(p.igst + delta * 0.18, 2))
        else:
            p = replace(p, cgst=round(p.cgst + delta * 0.09, 2), sgst=round(p.sgst + delta * 0.09, 2))
        books.append(b)
        portal.append(p)

    # 2b. Rule-sensitive case: ₹1,500 / ₹270 apart = ~0.6% -> over ₹10 but under 1%.
    #     ₹ only / ₹ AND % -> Mismatch (₹270 ITC shortfall at risk); % only / ₹ OR % -> Matched with Tolerance.
    b = taxed(SUP[3], "REW/T600", date(2024, 5, 16), 2_51_500,
              category="Rule-sensitive: Mismatch under ₹ only, Tolerance under % rules")
    books.append(b)
    portal.append(replace(b, taxable=2_50_000, igst=45_000))

    # 3. Mismatches (6)
    s = SUP
    b = taxed(s[4], "DOS-24-25-7001", date(2024, 4, 12), 50_000, category="Mismatch - taxable value")
    books.append(b); portal.append(replace(b, taxable=60_000, igst=10_800))
    b = taxed(s[2], "KEL/7002", date(2024, 4, 18), 80_000, 18, category="Mismatch - tax rate/amount")
    books.append(b); portal.append(replace(b, igst=9_600))
    b = taxed(s[5], "CTL/7003", date(2024, 5, 3), 42_000, category="Mismatch - invoice date")
    books.append(b); portal.append(replace(b, dt=date(2024, 5, 8)))
    b = taxed(s[0], "SGS/7004", date(2024, 5, 9), 1_20_000, head="IGST", category="Mismatch - tax head (IGST in books)")
    books.append(b); portal.append(taxed(s[0], "SGS/7004", date(2024, 5, 9), 1_20_000, category=b.category))
    b = taxed(s[6], "MTS/7005", date(2024, 4, 25), 15_000, head="IGST", category="Mismatch - tax head (IGST in books)")
    books.append(b); portal.append(taxed(s[6], "MTS/7005", date(2024, 4, 25), 15_000, category=b.category))
    b = taxed(s[7], "GL/7006", date(2024, 5, 14), 20_000, category="Mismatch - taxable beyond tolerance")
    books.append(b); portal.append(replace(b, taxable=20_250, igst=3_645))

    # 4. Probable matches (4) - typos in invoice numbers, amounts agree
    for sup, b_inv, p_inv, dt, tv in [
        (s[2], "KEL/56789", "KEL/56798", date(2024, 4, 22), 33_000),
        (s[5], "CTL/88901", "CTL/8891", date(2024, 5, 2), 18_500),
        (s[7], "GL/24-25/00917", "GL/917A", date(2024, 5, 20), 64_000),
        (s[9], "HPD/120110", "HPD/120101", date(2024, 4, 29), 9_999),
    ]:
        b = taxed(sup, b_inv, dt, tv, category="Probable Match")
        books.append(b); portal.append(replace(b, inv=p_inv))

    # 5. Missing in 2B (6) - incl. two invalid GSTINs in books
    for sup, inv, dt, tv in [(s[3], "REW/9001", date(2024, 4, 9), 72_000), (s[8], "BCC-9002", date(2024, 5, 11), 28_000),
                             (s[1], "PPI-9003", date(2024, 5, 27), 44_500), (s[9], "HPD/9004", date(2024, 5, 30), 12_000)]:
        books.append(taxed(sup, inv, dt, tv, category="Missing in 2B"))
    bad = taxed(s[4], "DOS-9005", date(2024, 4, 16), 16_000, category="Invalid GSTIN (checksum) -> Missing in 2B")
    wrong = bad.gstin[:-1] + ("A" if bad.gstin[-1] != "A" else "B")
    books.append(replace(bad, gstin=wrong))
    portal.append(replace(bad, category="Correct GSTIN on portal -> Missing in Books"))
    books.append(replace(taxed(s[8], "BCC-9006", date(2024, 5, 6), 8_000), gstin=s[8]["gstin"][:14],
                         category="Invalid GSTIN (length) -> Missing in 2B"))

    # 6. Missing in books (4 more + 1 above)
    for sup, inv, dt, tv, avl in [(s[1], "PPI-8001", date(2024, 4, 14), 5_600, "Y"), (s[3], "REW/8002", date(2024, 5, 4), 38_000, "Y"),
                                  (s[6], "MTS/8003", date(2024, 5, 19), 2_400, "N"), (s[9], "HPD/8004", date(2024, 4, 3), 51_000, "Y")]:
        p = taxed(sup, inv, dt, tv, category="Missing in Books" + (" (ITC not available)" if avl == "N" else ""))
        portal.append(replace(p, itcavl=avl))

    # 7. Credit / debit notes
    cn1 = taxed(s[2], "KEL/CN/011", date(2024, 4, 26), 5_000, doc="CN", category="Credit Note - exact")
    cn2 = taxed(s[0], "SGS-CN-012", date(2024, 5, 15), 8_000, doc="CN", category="Credit Note - exact")
    cn3 = taxed(s[5], "CTL/CN/013", date(2024, 5, 21), 3_500, doc="CN", category="Credit Note - missing in 2B")
    dn1 = taxed(s[4], "DOS-DN-014", date(2024, 5, 7), 2_000, doc="DN", category="Debit Note - exact")
    books += [cn1, cn2, cn3, dn1]
    portal += [replace(cn1, inv="KEL/CN/11"), cn2, dn1]
    return books, portal


def build_sales_data() -> tuple[list[Doc], list[Doc]]:
    """Return (sales books, GSTR-1) documents."""
    customers = [
        {"state": "27", "gstin": make_gstin("27", "AAFCA1111B"), "name": "Aarav Retail LLP", "prefix": "DT"},
        {"state": "29", "gstin": make_gstin("29", "AAGCB2222C"), "name": "Bangalore Mart Pvt Ltd", "prefix": "DT"},
        {"state": "32", "gstin": make_gstin("32", "AAHCK3333D"), "name": "Kochi Traders", "prefix": "DT"},
        {"state": "09", "gstin": make_gstin("09", "AAJCL4444E"), "name": "Lucknow Distributors", "prefix": "DT"},
    ]
    books: list[Doc] = []
    gstr1: list[Doc] = []
    for i in range(14):
        c = customers[i % 4]
        b = taxed(c, f"DT/24-25/{i + 1:04d}", date(2024, 5, 1 + i * 2), 15_000 + i * 4_250, category="Exact Match")
        books.append(b)
        gstr1.append(replace(b, inv=f"DT/{i + 1:04d}"))
    b = taxed(customers[1], "DT/24-25/0101", date(2024, 5, 10), 40_000, category="Matched with Tolerance")
    books.append(b); gstr1.append(replace(b, taxable=40_004, igst=7_200.72))
    b = taxed(customers[2], "DT/24-25/0102", date(2024, 5, 12), 22_000, category="Mismatch - taxable")
    books.append(b); gstr1.append(replace(b, taxable=24_000, igst=4_320))
    for n in (103, 104):
        books.append(taxed(customers[3], f"DT/24-25/{n:04d}", date(2024, 5, 28), 9_000 + n, category="Missing in GSTR-1"))
    gstr1.append(taxed(customers[0], "DT/24-25/0105", date(2024, 5, 30), 6_500, category="Missing in Books"))
    cn = taxed(customers[1], "DT/CN/0001", date(2024, 5, 25), 3_000, doc="CN", category="Credit Note - exact")
    books.append(cn); gstr1.append(cn)
    return books, gstr1


# --------------------------------------------------------------------------- #
# Writers
# --------------------------------------------------------------------------- #
BOOK_HEADERS = ["Supplier GSTIN", "Supplier Name", "Invoice No", "Invoice Date", "Invoice Type", "Taxable Value",
                "IGST", "CGST", "SGST", "Cess", "Total Invoice Value", "Place of Supply"]
TYPE_TEXT = {"INV": "B2B", "CN": "Credit Note", "DN": "Debit Note"}
BOLD = Font(bold=True)
HFILL = PatternFill("solid", start_color="DDEBF7", end_color="DDEBF7")


def _date_cell(d: date, i: int) -> object:
    """Mixed date representations, as commonly found in exported registers."""
    return [d, d.strftime("%d/%m/%Y"), d, d.strftime("%d-%b-%Y"), d, d.strftime("%Y-%m-%d")][i % 6]


def write_register(path: Path, docs: list[Doc], title: str, headers: list[str]) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "Register"
    ws.append([COMPANY_NAME])
    ws.append([title])
    ws.append([])
    ws.append(headers)
    for c in ws[4]:
        c.font, c.fill = BOLD, HFILL
    for i, d in enumerate(docs):
        ws.append([d.gstin, d.name, d.inv, _date_cell(d.dt, i), TYPE_TEXT[d.doc], d.taxable, d.igst, d.cgst, d.sgst,
                   d.cess, d.value, f"{d.pos}-{STATE_NAMES.get(d.pos, '')}"])
        cell = ws.cell(row=ws.max_row, column=4)
        if isinstance(cell.value, date):
            cell.number_format = "DD-MM-YYYY"
    n = ws.max_row
    ws.append(["", "Total", "", "", "", f"=SUM(F5:F{n})", f"=SUM(G5:G{n})", f"=SUM(H5:H{n})", f"=SUM(I5:I{n})",
               f"=SUM(J5:J{n})", f"=SUM(K5:K{n})", ""])
    for col, width in zip("ABCDEFGHIJKL", (18, 32, 20, 13, 13, 14, 12, 12, 12, 8, 16, 18)):
        ws.column_dimensions[col].width = width
    wb.save(path)


def write_tally_register(path: Path, docs: list[Doc], mode: str = "purchase") -> None:
    """Tally-style register export: title block, one ledger column per GST rate, round off, Dr suffixes."""
    rates = sorted({round((d.igst + d.cgst + d.sgst) / d.taxable * 100) for d in docs if d.taxable})
    trade = "Purchase" if mode == "purchase" else "Sales"
    io_word = "Input" if mode == "purchase" else "Output"
    headers = (["Date", "Particulars", "Voucher Type", "Voucher No."]
               + (["Supplier Invoice No.", "Supplier Invoice Date"] if mode == "purchase" else [])
               + ["GSTIN/UIN", "Gross Total"] + [f"{trade} @ {r}%" for r in rates]
               + [f"{io_word} CGST", f"{io_word} SGST", f"{io_word} IGST", "Round Off"])
    wb = Workbook()
    ws = wb.active
    ws.title = f"{trade} Register"
    ws.append([COMPANY_NAME])
    ws.append(["Andheri East, Mumbai - 400069"])
    ws.append([f"{trade} Register"])
    ws.append(["1-Apr-2024 to 31-May-2024"])
    ws.append([])
    ws.append(headers)
    for c in ws[6]:
        c.font = BOLD
    vch_type = {"INV": trade, "CN": "Debit Note" if mode == "purchase" else "Credit Note",
                "DN": "Credit Note" if mode == "purchase" else "Debit Note"}
    for n, d in enumerate(docs, start=1):
        value = d.value
        gross = float(round(value))
        rate = round((d.igst + d.cgst + d.sgst) / d.taxable * 100) if d.taxable else 0
        ledgers = [d.taxable if r == rate else None for r in rates]
        gross_cell: object = gross if n % 7 else f"{gross:,.2f} Dr"
        if mode == "purchase":  # our voucher number; the supplier's invoice number is a separate column
            row = [d.dt + timedelta(days=1), d.name, vch_type[d.doc], str(n), d.inv, d.dt]
        else:                   # sales: the voucher number is the invoice number
            row = [d.dt, d.name, vch_type[d.doc], d.inv]
        row += [d.gstin, gross_cell, *ledgers, d.cgst or None, d.sgst or None, d.igst or None,
                round(gross - value, 2) or None]
        ws.append(row)
        for col in (1, 6 if mode == "purchase" else None):
            if col:
                ws.cell(row=ws.max_row, column=col).number_format = "D-MMM-YY"
        if n == 10:  # a cancelled voucher, as Tally exports it
            ws.append([d.dt, "(cancelled)", trade, str(1000 + n)])
    ws.append([None, "Grand Total"])
    for i in range(1, len(headers) + 1):
        ws.column_dimensions[chr(64 + i) if i <= 26 else "A" + chr(64 + i - 26)].width = 15
    wb.save(path)


def write_gstr3b_json(path: Path, portal: list[Doc], period: str, adjust: dict[str, float]) -> None:
    """GSTR-3B JSON (Table 4): 4A(5) = ITC available in that month's 2B plus *adjust* per head."""
    heads = {"iamt": 0.0, "camt": 0.0, "samt": 0.0, "csamt": 0.0}
    for d in portal:
        if d.period == period and d.itcavl == "Y":
            sign = -1 if d.doc == "CN" else 1
            heads["iamt"] += sign * d.igst
            heads["camt"] += sign * d.cgst
            heads["samt"] += sign * d.sgst
            heads["csamt"] += sign * d.cess
    oth = {k: round(v + adjust.get(k, 0.0), 2) for k, v in heads.items()}
    zero = {"iamt": 0, "camt": 0, "samt": 0, "csamt": 0}
    payload = {"gstin": COMPANY_GSTIN, "ret_period": period, "itc_elg": {
        "itc_avl": [{"ty": "IMPG", **zero}, {"ty": "IMPS", **zero}, {"ty": "ISRC", **zero}, {"ty": "ISD", **zero},
                    {"ty": "OTH", **oth}],
        "itc_rev": [{"ty": "RUL", **zero}, {"ty": "OTH", **zero}],
        "itc_net": oth,
        "itc_inelg": [{"ty": "RUL", **zero}, {"ty": "OTH", **zero}]}}
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _group(docs: list[Doc]) -> dict[str, list[Doc]]:
    out: dict[str, list[Doc]] = {}
    for d in docs:
        out.setdefault(d.gstin, []).append(d)
    return out


def write_gstr2b_json(path: Path, docs: list[Doc], period: str) -> None:
    """Official GSTR-2B JSON layout (data.docdata.b2b[].inv[] / cdnr[].nt[])."""
    sel = [d for d in docs if d.period == period]
    b2b, cdnr = [], []
    for gstin, items in _group([d for d in sel if d.doc == "INV"]).items():
        b2b.append({"ctin": gstin, "trdnm": items[0].name, "supfildt": f"11-{period[:2]}-{period[2:]}", "supprd": period,
                    "inv": [{"inum": d.inv, "typ": "R", "dt": d.dt.strftime("%d-%m-%Y"), "val": d.value, "pos": d.pos,
                             "rev": "N", "itcavl": d.itcavl, "rsn": "", "diffprcnt": 1, "srctyp": "", "irn": "",
                             "irngendate": "", "txval": d.taxable, "igst": d.igst, "cgst": d.cgst, "sgst": d.sgst,
                             "cess": d.cess} for d in items]})
    for gstin, items in _group([d for d in sel if d.doc != "INV"]).items():
        cdnr.append({"ctin": gstin, "trdnm": items[0].name, "supfildt": f"11-{period[:2]}-{period[2:]}", "supprd": period,
                     "nt": [{"ntnum": d.inv, "typ": "C" if d.doc == "CN" else "D", "suptyp": "R",
                             "dt": d.dt.strftime("%d-%m-%Y"), "val": d.value, "pos": d.pos, "rev": "N",
                             "itcavl": d.itcavl, "rsn": "", "diffprcnt": 1, "srctyp": "", "irn": "", "irngendate": "",
                             "txval": d.taxable, "igst": d.igst, "cgst": d.cgst, "sgst": d.sgst, "cess": d.cess}
                            for d in items]})
    payload = {"chksum": "SAMPLE", "data": {"gstin": COMPANY_GSTIN, "rtnprd": period, "version": "GSTN",
                                            "gendt": f"14-{int(period[:2]) + 1:02d}-{period[2:]}",
                                            "docdata": {"b2b": b2b, "cdnr": cdnr}}}
    path.write_text(json.dumps(payload, indent=2), encoding="utf-8")


def _period_label(period: str) -> str:
    return date(int(period[2:]), int(period[:2]), 1).strftime("%b'%y")


def write_gstr2b_portal_excel(path: Path, docs: list[Doc]) -> None:
    """Portal-style GSTR-2B Excel: title rows, two-row merged headers, B2B and B2B-CDNR sheets."""
    wb = Workbook()
    readme = wb.active
    readme.title = "Read me"
    readme.append(["GSTR-2B - Auto-drafted ITC Statement (SAMPLE DATA)"])
    readme.append(["GSTIN", COMPANY_GSTIN])
    readme.append(["Legal name", COMPANY_NAME])

    ws = wb.create_sheet("B2B")
    ws.append(["Goods and Services Tax - GSTR-2B"])
    ws.append(["Taxable inward supplies received from registered persons"])
    ws.append([])
    ws.append([])
    ws.append(["GSTIN of supplier", "Trade/Legal name", "Invoice details", None, None, None, "Place of supply",
               "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount", None, None, None,
               "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability", "Reason",
               "Applicable % of Tax Rate", "Source", "IRN", "IRN Date"])
    ws.append([None, None, "Invoice number", "Invoice type", "Invoice Date", "Invoice Value(₹)", None, None, None,
               "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"])
    ws.merge_cells("C5:F5")
    ws.merge_cells("J5:M5")
    for col in "ABGHINOPQRSTU":
        ws.merge_cells(f"{col}5:{col}6")
    for d in [d for d in docs if d.doc == "INV"]:
        ws.append([d.gstin, d.name, d.inv, "Regular", d.dt.strftime("%d-%m-%Y"), d.value, STATE_NAMES[d.pos], "No",
                   d.taxable, d.igst, d.cgst, d.sgst, d.cess, _period_label(d.period),
                   f"11-{int(d.period[:2]) + 1:02d}-{d.period[2:]}", "Yes" if d.itcavl == "Y" else "No", "", "", "",
                   "", ""])

    ws2 = wb.create_sheet("B2B-CDNR")
    ws2.append(["Goods and Services Tax - GSTR-2B"])
    ws2.append(["Debit/Credit notes (Original)"])
    ws2.append([])
    ws2.append([])
    ws2.append(["GSTIN of supplier", "Trade/Legal name", "Credit note/Debit note details", None, None, None, None,
                "Place of supply", "Supply Attract Reverse Charge", "Taxable Value (₹)", "Tax Amount", None, None, None,
                "GSTR-1/IFF/GSTR-5 Period", "GSTR-1/IFF/GSTR-5 Filing Date", "ITC Availability", "Reason"])
    ws2.append([None, None, "Note number", "Note type", "Note Supply type", "Note date", "Note Value (₹)", None, None,
                None, "Integrated Tax(₹)", "Central Tax(₹)", "State/UT Tax(₹)", "Cess(₹)"])
    ws2.merge_cells("C5:G5")
    ws2.merge_cells("K5:N5")
    for col in "ABHIJOPQR":
        ws2.merge_cells(f"{col}5:{col}6")
    for d in [d for d in docs if d.doc != "INV"]:
        ws2.append([d.gstin, d.name, d.inv, "Credit Note" if d.doc == "CN" else "Debit Note", "Regular",
                    d.dt.strftime("%d-%m-%Y"), d.value, STATE_NAMES[d.pos], "No", d.taxable, d.igst, d.cgst, d.sgst,
                    d.cess, _period_label(d.period), "", "Yes" if d.itcavl == "Y" else "No", ""])
    for sheet in (ws, ws2):
        sheet["A1"].font = Font(bold=True, size=13)
        for row in sheet.iter_rows(min_row=5, max_row=6):
            for c in row:
                c.font, c.fill = BOLD, HFILL
                c.alignment = Alignment(horizontal="center", wrap_text=True)
        for i in range(1, 22):
            sheet.column_dimensions[chr(64 + i)].width = 16
    wb.save(path)


def write_gstr1_json(path: Path, docs: list[Doc], period: str) -> None:
    """GSTR-1 JSON layout (b2b[].inv[].itms[].itm_det / cdnr[].nt[])."""
    def items(d: Doc) -> list[dict]:
        rate = round((d.igst + d.cgst + d.sgst) / d.taxable * 100) if d.taxable else 0
        return [{"num": 1, "itm_det": {"rt": rate, "txval": d.taxable, "iamt": d.igst, "camt": d.cgst,
                                       "samt": d.sgst, "csamt": d.cess}}]
    b2b = [{"ctin": g, "inv": [{"inum": d.inv, "idt": d.dt.strftime("%d-%m-%Y"), "val": d.value, "pos": d.pos,
                                 "rchrg": "N", "inv_typ": "R", "itms": items(d)} for d in ds]}
           for g, ds in _group([d for d in docs if d.doc == "INV"]).items()]
    cdnr = [{"ctin": g, "nt": [{"ntty": "C" if d.doc == "CN" else "D", "nt_num": d.inv, "nt_dt": d.dt.strftime("%d-%m-%Y"),
                                "val": d.value, "pos": d.pos, "rchrg": "N", "inv_typ": "R", "itms": items(d)} for d in ds]}
            for g, ds in _group([d for d in docs if d.doc != "INV"]).items()]
    path.write_text(json.dumps({"gstin": COMPANY_GSTIN, "fp": period, "b2b": b2b, "cdnr": cdnr}, indent=2),
                    encoding="utf-8")


def write_templates() -> None:
    TEMPLATES.mkdir(exist_ok=True)
    example = ["27AAPFU0939F1ZV", "ABC Suppliers Pvt Ltd", "ABC/24-25/0001", "15-04-2024", "B2B", 100000, 0, 9000,
               9000, 0, 118000, "27-Maharashtra"]
    for fname, headers, rows in [
        ("purchase_register_template.xlsx", BOOK_HEADERS, [example, example[:2] + ["ABC/CN/01", "20-04-2024", "Credit Note",
                                                                                    5000, 0, 450, 450, 0, 5900, "27-Maharashtra"]]),
        ("gstr2b_template.xlsx", BOOK_HEADERS + ["Return Period", "ITC Availability"], [example + ["052024", "Y"]]),
        ("sales_register_template.xlsx", ["Customer GSTIN", "Customer Name"] + BOOK_HEADERS[2:], [example]),
    ]:
        wb = Workbook()
        ws = wb.active
        ws.title = "Data"
        ws.append(headers)
        for c in ws[1]:
            c.font, c.fill = BOLD, HFILL
        for r in rows:
            ws.append(r)
        for i in range(1, len(headers) + 1):
            ws.column_dimensions[chr(64 + i)].width = 18
        notes = wb.create_sheet("Instructions")
        for line in [
            "Keep one row per invoice / credit note / debit note.",
            "Invoice Type: B2B (invoice), Credit Note, Debit Note (or CDNR with a 'Note Type' column C/D).",
            "Dates: any common format (DD-MM-YYYY preferred). Amounts: numbers, commas allowed.",
            "Column names may differ - map them on the Column Mapping screen and save the mapping.",
        ]:
            notes.append([line])
        notes.column_dimensions["A"].width = 100
        wb.save(TEMPLATES / fname)


def main() -> None:
    OUT.mkdir(exist_ok=True)
    books, portal = build_purchase_data()
    write_register(OUT / "purchase_register.xlsx", books, "Purchase Register: 01-04-2024 to 31-05-2024", BOOK_HEADERS)
    write_gstr2b_json(OUT / "gstr2b_apr2024.json", portal, "042024")
    write_gstr2b_json(OUT / "gstr2b_may2024.json", portal, "052024")
    write_gstr2b_portal_excel(OUT / "gstr2b_portal_format.xlsx", portal)
    write_tally_register(OUT / "tally_purchase_register.xlsx", books, "purchase")
    write_gstr3b_json(OUT / "gstr3b_apr2024.json", portal, "042024", {"iamt": 5000.0})           # excess claim
    write_gstr3b_json(OUT / "gstr3b_may2024.json", portal, "052024", {"camt": -600.0, "samt": -600.0})  # short

    s_books, s_gstr1 = build_sales_data()
    sales_headers = ["Customer GSTIN", "Customer Name"] + BOOK_HEADERS[2:]
    write_register(OUT / "sales_register.xlsx", s_books, "Sales Register: May 2024", sales_headers)
    write_gstr1_json(OUT / "gstr1_may2024.json", s_gstr1, "052024")
    write_tally_register(OUT / "tally_sales_register.xlsx", s_books, "sales")

    with (OUT / "expected_categories.csv").open("w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(["dataset", "gstin", "invoice_no", "doc_type", "category"])
        for label, docs in (("books", books), ("gstr2b", portal), ("sales_books", s_books), ("gstr1", s_gstr1)):
            for d in docs:
                w.writerow([label, d.gstin, d.inv, d.doc, d.category])
    write_templates()
    print(f"Company GSTIN : {COMPANY_GSTIN}")
    print(f"Purchase books: {len(books)} rows | GSTR-2B: {len(portal)} documents")
    print(f"Sales books   : {len(s_books)} rows | GSTR-1 : {len(s_gstr1)} documents")
    print(f"Files written to {OUT} and {TEMPLATES}")


if __name__ == "__main__":
    main()
