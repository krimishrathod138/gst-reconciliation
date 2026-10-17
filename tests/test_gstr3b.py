"""Tests for the GSTR-3B vs 2B vs books comparison (modules.gstr3b)."""
from __future__ import annotations

import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from conftest import clean, doc
from modules import ReconError
from modules.gstr3b import combine_gstr3b, compare_gstr3b, parse_gstr3b_json, parse_gstr3b_table
from modules.matcher import reconcile


def g3b_json(period: str, i=0.0, c=0.0, s=0.0, cess=0.0, wrap=False) -> bytes:
    oth = {"ty": "OTH", "iamt": i, "camt": c, "samt": s, "csamt": cess}
    body = {"gstin": "27AABCD1234E1Z8", "ret_period": period,
            "itc_elg": {"itc_avl": [{"ty": "IMPG", "iamt": 99, "camt": 0, "samt": 0, "csamt": 0}, oth],
                        "itc_net": {"iamt": i + 99, "camt": c, "samt": s, "csamt": cess}}}
    return json.dumps({"data": body} if wrap else body).encode()


def test_parse_json_uses_table_4a5():
    row = parse_gstr3b_json(g3b_json("052024", i=1800, c=90, s=90)).iloc[0]
    assert row["period"] == "2024-05"
    assert (row["igst"], row["cgst"], row["sgst"], row["cess"]) == (1800, 90, 90, 0)   # imports (IMPG) excluded
    assert row["net_igst"] == 1899
    assert parse_gstr3b_json(g3b_json("052024", i=5, wrap=True)).iloc[0]["igst"] == 5


def test_parse_json_errors():
    with pytest.raises(ReconError, match="not recognised as GSTR-3B"):
        parse_gstr3b_json(b'{"b2b": []}')
    with pytest.raises(ReconError, match="Invalid"):
        parse_gstr3b_json(b"{bad")


def test_parse_table_and_combine():
    t = parse_gstr3b_table(pd.DataFrame({"Month": ["Apr-2024", "05-2024"], "Integrated Tax": ["1,000", 0],
                                         "Central Tax": [50, 60], "State/UT Tax": [50, 60]}))
    assert t["period"].tolist() == ["2024-04", "2024-05"] and t["igst"].tolist() == [1000, 0]
    assert t["cess"].tolist() == [0, 0]
    later = parse_gstr3b_json(g3b_json("052024", i=7))
    merged = combine_gstr3b([t, later])
    assert merged.set_index("period").loc["2024-05", "igst"] == 7   # later file replaces the month
    with pytest.raises(ReconError, match="needs columns"):
        parse_gstr3b_table(pd.DataFrame({"Month": ["Apr-2024"], "Tax": [1]}))


@pytest.fixture
def result():
    books = [doc("A", 10000, 1800, date="05-04-2024"), doc("B", 5000, 900, date="10-04-2024"),
             doc("CN1", 1000, 180, date="20-04-2024", typ="Credit Note"), doc("C", 2000, 0, 180, 180, date="02-05-2024")]
    portal = [doc("A", 10000, 1800, date="05-04-2024", period="042024"),
              doc("CN1", 1000, 180, date="20-04-2024", typ="CDNR", period="042024"),
              doc("B", 5000, 900, date="10-04-2024", period="052024"),          # supplier filed late
              doc("C", 2000, 0, 180, 180, date="02-05-2024", period="052024"),
              doc("N1", 1000, 180, date="03-05-2024", period="052024", itc_available="N")]
    return reconcile(clean(books), clean(portal, "portal"))


def test_compare_month_wise(result):
    g3b = combine_gstr3b([parse_gstr3b_json(g3b_json("042024", i=2620)),           # 1800 - 180 + 1000 excess
                          parse_gstr3b_json(g3b_json("052024", i=900, c=100, s=100))])  # 160 short
    out = compare_gstr3b(result, g3b).set_index("month")
    apr, may, tot = out.loc["2024-04"], out.loc["2024-05"], out.loc["TOTAL"]
    assert apr["books_total"] == 1800 + 900 - 180 and apr["gstr2b_total"] == 1620 and apr["gstr3b_total"] == 2620
    assert apr["diff_3b_2b_total"] == 1000 and apr["status"] == "Excess claimed"
    assert may["gstr2b_total"] == 900 + 360 and may["gstr2b_not_available"] == 180   # N1 not counted as available
    assert may["diff_3b_2b_cgst"] == -80 and may["status"] == "Short claimed"
    assert may["cumulative_3b_minus_2b"] == 1000 - 160
    assert tot["books_total"] == 2880 and tot["gstr2b_total"] == 2880 and tot["gstr3b_total"] == 3720
    assert tot["diff_books_2b"] == 0 and tot["status"] == "Excess claimed"


def test_compare_without_3b_and_month_filter(result):
    out = compare_gstr3b(result, None, months=["2024-05"])
    assert out["month"].tolist() == ["2024-05", "TOTAL"] and set(out["status"]) == {"3B not entered"}


def test_compare_rejects_sales():
    r = reconcile(clean([doc("A", 1, 0)]), clean([doc("A", 1, 0)], "portal"), mode="sales", portal_name="GSTR-1")
    with pytest.raises(ReconError, match="purchase"):
        compare_gstr3b(r, None)


def test_sample_3b_files(sample_dir):
    from modules.cleaner import clean_dataset
    from modules.loader import load_file, to_canonical
    books, _ = clean_dataset(to_canonical(load_file(sample_dir / "purchase_register.xlsx")))
    portal, _ = clean_dataset(pd.concat([to_canonical(load_file(sample_dir / f, dataset="portal"))
                                         for f in ("gstr2b_apr2024.json", "gstr2b_may2024.json")]), "portal")
    g3b = combine_gstr3b([parse_gstr3b_json((sample_dir / f).read_bytes())
                          for f in ("gstr3b_apr2024.json", "gstr3b_may2024.json")])
    out = compare_gstr3b(reconcile(books, portal), g3b).set_index("month")
    assert out.loc["2024-04", "diff_3b_2b_total"] == 5000 and out.loc["2024-04", "status"] == "Excess claimed"
    assert out.loc["2024-05", "diff_3b_2b_total"] == -1200 and out.loc["2024-05", "status"] == "Short claimed"


def test_excel_sheet(result):
    from modules.reports import build_excel_report, gstr3b_sheet
    g3b = parse_gstr3b_json(g3b_json("042024", i=2620))
    wb = load_workbook(io.BytesIO(build_excel_report(result, extra_sheets=gstr3b_sheet(compare_gstr3b(result, g3b)))))
    ws = wb["GSTR-3B vs 2B vs Books"]
    rows = [{h.value: c.value for h, c in zip(ws[1], r)} for r in ws.iter_rows(min_row=2)]
    assert rows[0]["Month"] == "Apr 2024" and rows[0]["3B − 2B (Total)"] == 1000 and rows[-1]["Month"] == "TOTAL"
    assert rows[0]["Status"] == "Excess claimed"
