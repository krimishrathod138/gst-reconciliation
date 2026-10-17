"""Tests for loader, reports and database (integration level, using the sample data)."""
from __future__ import annotations

import io
import json

import pandas as pd
import pytest
from openpyxl import load_workbook

from config import INR_FORMAT, STATUS_EXACT, STATUS_MISMATCH, STATUS_PROBABLE
from conftest import VALID_GSTIN, clean, doc
from modules import ReconError
from modules.cleaner import clean_dataset
from modules.database import Database
from modules.formatting import format_inr
from modules.loader import (
    apply_mapping, detect_header, load_file, missing_required, parse_portal_json, suggest_mapping, to_canonical,
)
from modules.matcher import reconcile
from modules.reports import build_excel_report, supplier_email


# ----------------------------------------------------------------- loader
def test_suggest_mapping_portal_headers():
    cols = ["GSTIN of supplier", "Trade/Legal name", "Invoice number", "Invoice type", "Invoice Date",
            "Invoice Value(₹)", "Place of supply", "Taxable Value (₹)", "Integrated Tax(₹)", "Central Tax(₹)",
            "State/UT Tax(₹)", "Cess(₹)", "GSTR-1/IFF/GSTR-5 Period", "ITC Availability"]
    m = suggest_mapping(cols)
    assert m["gstin"] == "GSTIN of supplier" and m["invoice_no"] == "Invoice number"
    assert m["sgst"] == "State/UT Tax(₹)" and m["return_period"] == "GSTR-1/IFF/GSTR-5 Period"
    assert m["invoice_value"] == "Invoice Value(₹)" and not missing_required(m)


def test_suggest_mapping_prefers_specific_date_column():
    m = suggest_mapping(["Date", "Supplier Invoice Date", "Party GSTIN", "Bill No", "Taxable Amount"])
    assert m["invoice_date"] == "Supplier Invoice Date"
    assert m["invoice_no"] == "Bill No" and m["taxable_value"] == "Taxable Amount"


def test_apply_mapping_requires_fields():
    with pytest.raises(ReconError, match="required"):
        apply_mapping(pd.DataFrame({"x": [1]}), {"gstin": "x"})


def test_detect_two_row_header():
    raw = pd.DataFrame([
        ["Title", None, None, None],
        ["GSTIN of supplier", "Invoice details", None, "Taxable Value"],
        [None, "Invoice number", "Invoice Date", None],
        [VALID_GSTIN, "A1", "01-04-2024", 100],
    ])
    assert detect_header(raw) == (1, 2)


def test_load_sample_register(sample_dir):
    lf = load_file(sample_dir / "purchase_register.xlsx")
    assert lf.kind == "table" and "Supplier GSTIN" in lf.frame.columns
    books, rep = clean_dataset(to_canonical(lf))
    assert len(books) == 53 and rep.dropped_blank == 1
    assert books["invoice_date"].notna().all()


def test_csv_with_title_rows():
    text = "My Company\nPurchase register\nGSTIN,Invoice No,Invoice Date,Taxable Value,IGST\n" \
           f"{VALID_GSTIN},A-1,01/04/2024,\"1,000.00\",180\n"
    lf = load_file(text.encode(), name="books.csv")
    books, _ = clean_dataset(to_canonical(lf))
    assert books.loc[0, "taxable_value"] == 1000 and books.loc[0, "inv_norm"] == "A1"


def test_portal_json_and_excel_agree(sample_dir):
    js = pd.concat([to_canonical(load_file(sample_dir / f, dataset="portal"))
                    for f in ("gstr2b_apr2024.json", "gstr2b_may2024.json")], ignore_index=True)
    xl = load_file(sample_dir / "gstr2b_portal_format.xlsx", dataset="portal")
    assert xl.kind == "portal_excel"
    a, _ = clean_dataset(js, "portal")
    b, _ = clean_dataset(to_canonical(xl), "portal")
    key = ["gstin", "inv_norm", "doc_type", "taxable_value", "total_tax", "period", "itc_available"]
    assert a[key].sort_values(key).reset_index(drop=True).equals(b[key].sort_values(key).reset_index(drop=True))
    assert set(a["doc_type"]) == {"INV", "CN", "DN"}


def test_gstr1_itemwise_json():
    payload = {"gstin": VALID_GSTIN, "fp": "052024", "b2b": [{"ctin": VALID_GSTIN, "inv": [
        {"inum": "S1", "idt": "02-05-2024", "val": 236, "pos": "27", "itms": [
            {"num": 1, "itm_det": {"txval": 100, "camt": 9, "samt": 9, "csamt": 0}},
            {"num": 2, "itm_det": {"txval": 100, "camt": 9, "samt": 9, "csamt": 0}}]}]}],
        "cdnr": [{"ctin": VALID_GSTIN, "nt": [{"ntty": "C", "nt_num": "CN1", "nt_dt": "03-05-2024", "val": 59,
                                               "itms": [{"itm_det": {"txval": 50, "iamt": 9}}]}]}]}
    frame, meta = parse_portal_json(json.dumps(payload).encode())
    out, _ = clean_dataset(frame, "portal")
    assert meta["return_period"] == "052024"
    assert out["taxable_value"].tolist() == [200, 50] and out["cgst"].tolist() == [18, 0]
    assert out["doc_type"].tolist() == ["INV", "CN"] and out["period"].tolist() == ["2024-05", "2024-05"]


def test_unreconciled_2b_sections_are_reported():
    payload = {"data": {"gstin": VALID_GSTIN, "rtnprd": "052024", "docdata": {
        "b2b": [{"ctin": VALID_GSTIN, "inv": [{"inum": "A1", "dt": "01-05-2024", "txval": 100, "igst": 18}]}],
        "isd": [{"ctin": VALID_GSTIN, "doclist": [{"docnum": "I1"}, {"docnum": "I2"}]}],
        "impg": [{"refdt": "01-05-2024", "portcode": "INNSA1"}]}}}
    lf = load_file(json.dumps(payload).encode(), name="2b.json", dataset="portal")
    assert lf.meta["skipped_sections"] == {"ISD": 2, "IMPG": 1}
    assert any("ISD (2)" in n and "IMPG (1)" in n for n in lf.notes)
    assert len(lf.frame) == 1


def test_bad_json_message():
    with pytest.raises(ReconError, match="not recognised"):
        parse_portal_json(b'{"foo": 1}')
    with pytest.raises(ReconError, match="Invalid JSON"):
        parse_portal_json(b"{oops")


def test_unsupported_extension():
    with pytest.raises(ReconError, match="Unsupported"):
        load_file(b"abc", name="file.pdf")


# ----------------------------------------------------------------- full sample reconciliation
@pytest.fixture(scope="module")
def sample_result():
    from pathlib import Path
    root = Path(__file__).resolve().parents[1] / "sample_data"
    books, _ = clean_dataset(to_canonical(load_file(root / "purchase_register.xlsx")))
    portal = pd.concat([to_canonical(load_file(root / f, dataset="portal"))
                        for f in ("gstr2b_apr2024.json", "gstr2b_may2024.json")], ignore_index=True)
    portal, _ = clean_dataset(portal, "portal")
    return reconcile(books, portal)


def test_sample_covers_every_category(sample_result):
    counts = sample_result.recon["status"].value_counts().to_dict()
    assert counts == {STATUS_EXACT: 29, "Missing in GSTR-2B": 8, STATUS_MISMATCH: 7, "Missing in Books": 5,
                      "Matched with Tolerance": 5, STATUS_PROBABLE: 4}
    assert sample_result.recon["cross_period"].sum() == 3
    s = sample_result.summary()
    assert s["invalid_gstin_books"] == 2 and s["duplicates_books"] == 2
    assert s["itc_at_risk"] == 39360


@pytest.mark.parametrize("mode, status, at_risk", [
    ("absolute", STATUS_MISMATCH, 39360), ("both", STATUS_MISMATCH, 39360),
    ("percentage", "Matched with Tolerance", 39090), ("either", "Matched with Tolerance", 39090),
])
def test_sample_rule_sensitive_invoice(sample_result, mode, status, at_risk):
    """REW/T600 differs by ₹1,500 taxable / ₹270 IGST (~0.6%): over ₹10 but under 1%."""
    from config import ReconSettings
    r = reconcile(sample_result.books, sample_result.portal, ReconSettings(tolerance_mode=mode))
    row = r.recon.set_index("invoice_no_books").loc["REW/T600"]
    assert row["status"] == status
    assert row["diff_taxable_value"] == 1500 and row["diff_igst"] == 270
    assert r.summary()["itc_at_risk"] == at_risk


def test_excel_report_structure(sample_result, tmp_path):
    path = tmp_path / "report.xlsx"
    build_excel_report(sample_result, "Demo", "27AABCD1234E1Z8", "Apr-May 2024", path=path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Matched", "Matched with Tolerance", "Mismatched", "Probable Matches",
                             "Missing in 2B", "Missing in Books", "Supplier-wise Summary", "Duplicates",
                             "Invalid GSTINs"]
    ws = wb["Mismatched"]
    assert ws.freeze_panes == "C2" and ws.max_row == 8  # header + 7 mismatches
    headers = [c.value for c in ws[1]]
    diff_col = headers.index("Diff Taxable value") + 1
    assert ws.cell(1, diff_col).fill.start_color.rgb.endswith("C00000")
    assert ws.cell(2, headers.index("Taxable value (Books)") + 1).number_format == INR_FORMAT
    assert ws.cell(2, headers.index("Invoice Date (Books)") + 1).number_format == "dd-mm-yyyy"
    assert len(ws.conditional_formatting) > 0
    assert wb["Duplicates"].max_row == 3 and wb["Invalid GSTINs"].max_row == 3
    summary = {row[0]: row[1] for row in wb["Summary"].iter_rows(max_row=8, max_col=2, values_only=True) if row[0]}
    assert summary["Tolerance rule"] == "Within ±₹10 per amount field (₹ only); dates must match"


def test_summary_sheet_shows_rule_of_the_run(tmp_path):
    from config import ReconSettings
    r = reconcile(clean([doc("A1", 100, 18)]), clean([doc("A1", 100, 18)], "portal"),
                  ReconSettings(tolerance_mode="either", pct_tolerance=2))
    wb = load_workbook(io.BytesIO(build_excel_report(r)))
    values = [c for row in wb["Summary"].iter_rows(max_row=8, values_only=True) for c in row if c]
    assert "Within ±₹10 or ±2% per amount field (₹ OR %); dates must match" in values


def test_supplier_email_lists_missing(sample_result):
    rec = sample_result.recon
    gstin = rec.loc[rec["status"] == "Missing in GSTR-2B", "gstin"].iloc[0]
    text = supplier_email(sample_result, gstin, "Demo Traders", period_label="Apr 2024")
    assert "NOT reflected" in text and gstin in text and "Subject:" in text
    matched_only = rec.loc[rec["status"] == STATUS_EXACT, "gstin"]
    clean_gstins = set(matched_only) - set(rec.loc[rec["status"].isin(["Missing in GSTR-2B", STATUS_MISMATCH]), "gstin"])
    for g in clean_gstins:
        assert supplier_email(sample_result, g) == ""


@pytest.mark.parametrize("value, expected", [
    (0, "₹0.00"), (999.5, "₹999.50"), (1000, "₹1,000.00"), (100000, "₹1,00,000.00"),
    (12345678.9, "₹1,23,45,678.90"), (-250, "-₹250.00"), (None, ""),
])
def test_format_inr(value, expected):
    assert format_inr(value) == expected


# ----------------------------------------------------------------- database
def test_database_roundtrip(tmp_path, sample_result):
    db = Database(f"sqlite:///{(tmp_path / 'test.db').as_posix()}")
    run_id = db.save_run(sample_result, "Apr-May 2024", "Demo", books_files=["a.xlsx"], portal_files=["b.json"])
    runs = db.list_runs()
    assert runs.loc[0, "run_id"] == run_id and runs.loc[0, "period"] == "Apr-May 2024"
    loaded, decisions, meta = db.load_run(run_id)
    assert decisions.empty and meta["books_files"] == ["a.xlsx"]
    assert loaded.summary() == sample_result.summary()
    mid = loaded.recon.loc[loaded.recon["status"] == STATUS_PROBABLE, "match_id"].iloc[0]
    assert db.save_decisions(run_id, pd.DataFrame({"match_id": [mid], "decision": ["Accepted"], "remarks": ["ok"]}),
                             "asha") == 1
    assert db.save_decisions(run_id, pd.DataFrame({"match_id": [mid], "decision": ["Accepted"], "remarks": ["ok"]}),
                             "asha") == 0  # unchanged -> no audit row
    db.save_decisions(run_id, pd.DataFrame({"match_id": [mid], "decision": ["Rejected"], "remarks": ["no"]}), "ravi")
    db.save_decisions(run_id, pd.DataFrame({"match_id": [mid], "decision": [""], "remarks": ["no - wrong party"]}),
                      "ravi")  # remark-only change keeps the decision
    dec = db.load_decisions(run_id)
    rec = dec.iloc[0]
    assert (rec["match_id"], rec["decision"], rec["remarks"], rec["updated_by"]) == (mid, "Rejected", "no - wrong party", "ravi")
    audit = db.load_audit(run_id)
    assert audit["action"].tolist() == ["Remark", "Rejected", "Accepted"]            # newest first
    assert audit["username"].tolist() == ["ravi", "ravi", "asha"]
    assert audit.iloc[1][["old_decision", "new_decision"]].tolist() == ["Accepted", "Rejected"]
    assert audit["at"].notna().all()
    view = loaded.with_decisions(dec)
    assert f"{mid}-B" in set(view.recon["match_id"])
    build_excel_report(view)  # a reloaded + reviewed result still produces a report
    db.delete_run(run_id)
    assert db.list_runs().empty


def test_reviewer_shown_on_probable_sheet(sample_result):
    mid = sample_result.recon.loc[sample_result.recon["status"] == STATUS_PROBABLE, "match_id"].iloc[0]
    dec = pd.DataFrame({"match_id": [mid], "decision": ["Pending"], "remarks": ["checking"], "updated_by": ["asha"],
                        "updated_at": [pd.Timestamp("2024-06-01 10:30")]})
    ws = load_workbook(io.BytesIO(build_excel_report(sample_result.with_decisions(dec))))["Probable Matches"]
    rows = [{h.value: c.value for h, c in zip(ws[1], r)} for r in ws.iter_rows(min_row=2)]
    row = next(r for r in rows if r["Match ID"] == mid)
    assert row["Reviewed By"] == "asha" and row["Reviewed At"].strftime("%d-%m-%Y %H:%M") == "01-06-2024 10:30"


def test_database_migrates_old_schema(tmp_path):
    import sqlite3
    path = tmp_path / "old.db"
    con = sqlite3.connect(path)
    con.execute("CREATE TABLE recon_runs (id INTEGER PRIMARY KEY, created_at DATETIME, mode VARCHAR(20), "
                "portal_name VARCHAR(20), period_label VARCHAR(100), company_name VARCHAR(200), "
                "company_gstin VARCHAR(15), books_files TEXT, portal_files TEXT, settings_json TEXT, "
                "summary_json TEXT, recon_blob BLOB, books_blob BLOB, portal_blob BLOB)")
    con.execute("INSERT INTO recon_runs (id, mode, portal_name, period_label) VALUES (1, 'purchase', 'GSTR-2B', 'old')")
    con.commit()
    con.close()
    db = Database(f"sqlite:///{path.as_posix()}")
    cols = {r[1] for r in sqlite3.connect(path).execute("pragma table_info(recon_runs)")}
    assert {"fy", "from_month", "to_month", "gstr3b_json", "created_by"} <= cols
    assert db.list_runs().loc[0, "period"] == "old"
    assert db.list_vendors().empty and db.load_audit(1).empty


def test_mapping_profiles(tmp_path):
    db = Database(f"sqlite:///{(tmp_path / 'm.db').as_posix()}")
    db.save_mapping("Tally export", "books", {"gstin": "Party GSTIN", "invoice_no": "Bill No", "cess": None})
    assert db.list_mappings("books") == ["Tally export"] and db.list_mappings("portal") == []
    assert db.load_mapping("Tally export") == {"gstin": "Party GSTIN", "invoice_no": "Bill No"}
    db.save_mapping("Tally export", "books", {"gstin": "GSTIN"})
    assert db.load_mapping("Tally export") == {"gstin": "GSTIN"}
    with pytest.raises(ReconError):
        db.save_mapping("  ", "books", {})
    db.delete_mapping("Tally export")
    assert db.list_mappings() == []


def test_indian_format_chosen_by_sign_and_size():
    from config import INR_FORMAT_ARAB, INR_FORMAT_KHARAB, INR_FORMAT_NEG
    vals = [500.0, 2_50_000.0, 5_00_00_00_000.0, 5_00_00_00_00_000.0]
    books = clean([doc(f"A{i}", v) for i, v in enumerate(vals)])
    portal = clean([doc(f"A{i}", v * 2) for i, v in enumerate(vals)], "portal")   # diffs are negative
    wb = load_workbook(io.BytesIO(build_excel_report(reconcile(books, portal))))
    ws = wb["Mismatched"]
    headers = [c.value for c in ws[1]]
    diff_col = headers.index("Diff Taxable value") + 1
    books_col = headers.index("Taxable value (Books)") + 1
    by_value = {ws.cell(r, books_col).value: (ws.cell(r, books_col).number_format, ws.cell(r, diff_col).number_format)
                for r in range(2, ws.max_row + 1)}
    assert by_value[500.0] == (INR_FORMAT, INR_FORMAT_NEG)
    assert by_value[2_50_000.0] == (INR_FORMAT, INR_FORMAT_NEG)
    assert by_value[5_00_00_00_000.0] == (INR_FORMAT_ARAB, INR_FORMAT_ARAB)
    assert by_value[5_00_00_00_00_000.0] == (INR_FORMAT_KHARAB, INR_FORMAT_KHARAB)
    assert ws.cell(2, diff_col).value < 0


def test_report_text_is_escaped_and_never_a_formula():
    books = clean([doc("X\x01Y&1", 1000, 180, name='A & B <Traders> "Ltd"')])
    portal = clean([doc("XY&1", 1000, 180, date="17-04-2024")], "portal")
    r = reconcile(books, portal)
    r.recon["remarks"] = "=SUM(A1:A9) & check"
    ws = load_workbook(io.BytesIO(build_excel_report(r)))["Mismatched"]
    row = {h.value: c for h, c in zip(ws[1], ws[2])}
    assert row["Supplier Name"].value == 'A & B <Traders> "Ltd"'
    assert row["Invoice No (Books)"].value == "XY&1"          # control character removed
    assert row["User Remarks"].value == "=SUM(A1:A9) & check" and row["User Remarks"].data_type == "s"
    assert row["Invoice Date (Books)"].value.strftime("%d-%m-%Y") == "15-04-2024"
    assert row["Date Diff (Days)"].value == -2


def test_large_report_speed_and_integrity():
    import time
    n = 20_000
    books = clean([doc(f"INV/{i:06d}", 1000 + i, 180) for i in range(n)])
    portal = clean([doc(f"INV/{i:06d}", 1000 + i, 180) for i in range(n)], "portal")
    r = reconcile(books, portal)
    t = time.perf_counter()
    data = build_excel_report(r)
    # Generous guard against a real regression only: wall-clock limits are flaky on a busy machine
    # (20k rows take ~5 s idle; the old per-cell writer needed ~15 s).
    assert time.perf_counter() - t < 180
    ws = load_workbook(io.BytesIO(data), read_only=True)["Matched"]
    rows = list(ws.iter_rows(min_row=n + 1, max_row=n + 1, values_only=True))
    assert ws.max_row == n + 1 and rows[0][5] == f"INV/{n - 1:06d}"


def test_small_reconcile_report_in_memory():
    r = reconcile(clean([doc("A1", 100, 18)]), clean([doc("A1", 100, 18)], "portal"))
    data = build_excel_report(r)
    wb = load_workbook(io.BytesIO(data))
    assert wb["Matched"].max_row == 2
